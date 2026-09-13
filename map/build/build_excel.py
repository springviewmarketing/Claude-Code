import json, os, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

practices=json.load(open("out/practices.json"))
companies=json.load(open("out/companies_classified.json"))
enr={}
if os.path.exists("out/enrichment.json"):
    for e in json.load(open("out/enrichment.json")): enr[e["OrgId"]]=e
import re as _re2
PLACEHOLDER=_re2.compile(r'@(domain|example|yourdomain|yoursite|mysite|email|company|test|sample|acme|website)\.',_re2.I)
JUNK=_re2.compile(r'(sentry|wixpress|godaddy|squarespace|\.png|\.jpg|\.svg|u002|sentry\.io|noreply@|no-reply@)',_re2.I)
def clean_email(em, site):
    if not em: return ""
    if PLACEHOLDER.search(em) or JUNK.search(em): return ""
    if em.count("@")!=1 or len(em)<6: return ""
    return em
dropped=0
for p in practices:
    e=enr.get(p["OrgId"],{})
    p["Website"]=e.get("Website","")
    raw=e.get("Email","")
    p["Email"]=clean_email(raw,p["Website"])
    if raw and not p["Email"]: dropped+=1
print("placeholder/junk emails dropped:",dropped)

NAVY="1F3864"; LBLUE="D9E2F3"; GREY="F2F2F2"; GREEN="C6E0B4"; AMBER="FFE699"
H=Font(bold=True,color="FFFFFF",size=11); HF=PatternFill("solid",fgColor=NAVY)
TITLE=Font(bold=True,size=15,color=NAVY); SUB=Font(bold=True,size=12,color=NAVY)
thin=Side(style="thin",color="BFBFBF"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook(); wb.remove(wb.active)

def sheet(name,rows,cols,widths=None,freeze="A2"):
    ws=wb.create_sheet(name)
    ws.append(cols)
    for c in range(1,len(cols)+1):
        cell=ws.cell(1,c); cell.font=H; cell.fill=HF
        cell.alignment=Alignment(vertical="center",wrap_text=True)
    ws.row_dimensions[1].height=30
    for r in rows: ws.append([r.get(c,"") for c in cols])
    ws.freeze_panes=freeze
    ws.auto_filter.ref=f"A1:{get_column_letter(len(cols))}{len(rows)+1}"
    for i,c in enumerate(cols,1):
        w=(widths or {}).get(c, min(max(12,len(c)+4,*(len(str(r.get(c,"")))+2 for r in rows[:300])) if rows else 16,46))
        ws.column_dimensions[get_column_letter(i)].width=w
    return ws

# ---------- HEADLINE ----------
ws=wb.create_sheet("Headline Figures")
ws.column_dimensions["A"].width=62; ws.column_dimensions["B"].width=16; ws.column_dimensions["C"].width=13
ws.column_dimensions["D"].width=64
r=1
def put(a,b="",c="",d="",font=None,fill=None):
    global r
    ws.cell(r,1,a); ws.cell(r,2,b); ws.cell(r,3,c); ws.cell(r,4,d)
    for col in range(1,5):
        cell=ws.cell(r,col)
        if font: cell.font=font
        if fill: cell.fill=PatternFill("solid",fgColor=fill)
        cell.alignment=Alignment(vertical="center",wrap_text=(col==4))
    if isinstance(b,(int,float)): ws.cell(r,2).number_format="#,##0"
    if isinstance(c,str) and c.endswith("%"): pass
    r+=1
def blank(n=1):
    global r
    r+=n

put("UK OPTICIANS: MARKET STRUCTURE & GROUP OWNERSHIP",font=TITLE); blank()
put(f"Prepared {datetime.date.today():%d %B %Y}  |  Companies House snapshot 01 Sep 2026  |  PSC snapshot 12 Sep 2026  |  NHS ODS live")
blank()

tot=len(practices)
cat=collections.Counter(p["OwnershipCategory"] for p in practices)
grp=collections.Counter(p["OwnershipGroup"] for p in practices)
indep=sum(v for k,v in cat.items() if k.startswith("Independent"))
chain=cat["Multiple / national chain"]; hak=cat["Hakim Group (consolidator)"]
grpowned=cat["Group-owned (6+ practices)"]
standalone=sum(1 for p in practices if p["StandalonePractice"]=="Yes")

put("SECTION 1 — PRACTICE LEVEL (England & Wales, NHS-registered, active)",font=SUB,fill=LBLUE)
put("Metric","Count","Share","Notes",font=Font(bold=True),fill=GREY)
def pct(n): return f"{n/tot*100:.1f}%"
put("Total optical practices",tot,"100%","NHS ODS 'Optical Site' (RO167), active records")
put("  Independently owned (all sizes)",indep,pct(indep),"Standalone + owner-operated groups of 2-5")
put("    of which: standalone single practice",cat['Independent - standalone'],pct(cat['Independent - standalone']),"Only practice under its owner")
put("    of which: independent small group (2-5)",indep-cat['Independent - standalone'],pct(indep-cat['Independent - standalone']),"Owner-operated mini-chains")
put("  Hakim Group (trades as independent)",hak,pct(hak),"Owned via HO2 Management Ltd, Darwen BB3 1AE")
put("  Other group / consolidator (6+)",grpowned,pct(grpowned),"Regional chains and roll-ups")
put("  National multiples / chains",chain,pct(chain),"Specsavers, Boots, Vision Express, Asda, Tesco etc.")
blank()
put("KEY FINDING","","", "")
put(f"Of {tot:,} NHS-listed practices in England & Wales, {indep:,} ({pct(indep)}) are genuinely independent. "
    f"A further {hak:,} ({pct(hak)}) trade under independent-looking local names but are owned by Hakim Group. "
    f"{chain+grpowned:,} ({pct(chain+grpowned)}) are chains or groups.",fill=GREEN)
blank(2)

put("SECTION 2 — PRACTICES BY OWNER / BRAND (England & Wales)",font=SUB,fill=LBLUE)
put("Owner / brand","Practices","Share","Category",font=Font(bold=True),fill=GREY)
CATOF={}
for p in practices: CATOF.setdefault(p["OwnershipGroup"],set()).add(p["OwnershipCategory"])
for k,v in grp.most_common():
    cats=CATOF.get(k,set())
    lbl=("Independent - owner-operated group of 2-5" if k=="Independent mini-group"
         else (list(cats)[0] if len(cats)==1 else "; ".join(sorted(cats))[:60]))
    put(k,v,pct(v),lbl)
blank(2)

put("SECTION 3 — HAKIM GROUP FOOTPRINT",font=SUB,fill=LBLUE)
put("Metric","Count","","Notes",font=Font(bold=True),fill=GREY)
hakco=sum(1 for c in companies if c["OwnershipGroup"]=="Hakim Group")
hakall=len(json.load(open("out/hakim_all_companies.json")))
put("NHS practices in England & Wales",hak,"","Practice-level, the most reliable count")
put("UK optical companies owned by Hakim",hakco,"","Companies House + PSC, incl. Scotland & N. Ireland")
put("All UK companies under Hakim control",hallk:=hakall,"","Includes audiology (Amplify Hearing), property, dormant")
put("Hakim's own public claim","500+","","UK AND Ireland, incl. audiology - consistent with the above")
blank()
put("How Hakim ownership was identified:", font=Font(bold=True))
put("Hakim practices keep their original local trading names, so name matching finds almost nothing. "
    "They are identified instead through Companies House Persons with Significant Control (PSC) data: the "
    "controlling entity is HO2 Management Ltd (and related vehicles) at Unit 317 India Mill Business Centre, "
    "Darwen BB3 1AE. Every company with an active Hakim-linked PSC was captured.",fill=AMBER)
blank(2)

put("SECTION 4 — COMPANY LEVEL (whole UK, incl. Scotland & N. Ireland)",font=SUB,fill=LBLUE)
put("Metric","Count","Share","Notes",font=Font(bold=True),fill=GREY)
ct=len(companies); cc=collections.Counter(c["OwnershipCategory"] for c in companies)
cind=sum(v for k,v in cc.items() if k.startswith("Independent"))
cch=cc["Multiple / national chain"]; chak=cc["Hakim Group (consolidator)"]
cgrp=sum(v for k,v in cc.items() if k.startswith("Group-owned"))
def cpct(n): return f"{n/ct*100:.1f}%"
put("Optician companies on Companies House",ct,"100%","SIC 47782 + optical businesses under health SICs")
put("  Independently owned",cind,cpct(cind),"")
put("  Hakim Group",chak,cpct(chak),"")
put("  Other group / consolidator",cgrp,cpct(cgrp),"")
put("  National multiples",cch,cpct(cch),"Mostly Specsavers JV store companies")
blank()
nat=collections.Counter("Scotland" if c["CompanyNumber"].startswith("SC") else
                        "Northern Ireland" if c["CompanyNumber"].startswith("NI") else
                        "England & Wales" for c in companies)
put("By nation of registration",font=Font(bold=True))
for k,v in nat.most_common(): put(f"  {k}",v,cpct(v),"")
blank(2)
put("SECTION 5 — CONTACT DATA COVERAGE",font=SUB,fill=LBLUE)
put("Field","Count","Share","Notes",font=Font(bold=True),fill=GREY)
wsite=sum(1 for p in practices if p["Website"]); mail=sum(1 for p in practices if p["Email"])
iw=sum(1 for p in practices if p["Website"] and p["OwnershipCategory"].startswith("Independent"))
im=sum(1 for p in practices if p["Email"] and p["OwnershipCategory"].startswith("Independent"))
ni=sum(1 for p in practices if p["OwnershipCategory"].startswith("Independent"))
put("Practices with a verified website",wsite,pct(wsite),"Includes chain sites; independents verified individually")
put("Practices with an email address",mail,pct(mail),"Scraped from the verified practice website only")
put("  Independent practices with website",iw,f"{iw/ni*100:.1f}%",f"Share of the {ni:,} independent practices")
put("  Independent practices with email",im,f"{im/ni*100:.1f}%",f"Share of the {ni:,} independent practices")
put("A blank is 'not confidently found', NOT 'does not exist'. Matches were only accepted where the domain "
    "resolved without cross-domain redirect and the page content confirmed an opticians, so the list is "
    "conservative by design and safe to use for outreach.",fill=AMBER)
ws.freeze_panes="A3"

# ---------- PRACTICES ----------
PC=["PracticeName","Address1","Address2","Town","County","Postcode","Country",
 "OwnershipCategory","OwnershipGroup","UltimateOwner","StandalonePractice","PracticesUnderSameOwner",
 "Website","Email","OperatingHQ","CompanyName","CompanyNumber","CompanyStatus","LegalForm","OrgId"]
W={"PracticeName":38,"Address1":30,"Address2":22,"OwnershipCategory":30,"OwnershipGroup":26,
   "UltimateOwner":38,"Website":38,"Email":34,"OperatingHQ":34,"CompanyName":34}
sheet("All Practices (E&W)",sorted(practices,key=lambda x:(x["OwnershipCategory"],x["PracticeName"])),PC,W)
ind=[p for p in practices if p["OwnershipCategory"].startswith("Independent")]
sheet("Independent Practices",sorted(ind,key=lambda x:(x["County"] or x["Town"],x["PracticeName"])),PC,W)
hk=[p for p in practices if p["OwnershipGroup"]=="Hakim Group"]
sheet("Hakim Group Practices",sorted(hk,key=lambda x:x["PracticeName"]),PC,W)

# ---------- LEAGUE TABLE ----------
lt=collections.Counter()
meta={}
for p in practices:
    k=p["UltimateOwner"] or p["OwnershipGroup"]
    lt[k]+=1; meta[k]=p["OwnershipCategory"]
rows=[{"Owner":k,"Practices":v,"Category":meta[k]} for k,v in lt.most_common() if v>=2]
sheet("Owner League Table",rows,["Owner","Practices","Category"],{"Owner":52,"Category":32})

# ---------- COMPANIES ----------
CC=["CompanyName","CompanyNumber","Postcode","Town","County","Address1","Address2",
 "CompanyStatus","Incorporated","OwnershipCategory","OwnershipGroup","UltimateOwner",
 "SIC1","ConfidenceTier","CompaniesHouseURL"]
sheet("UK Companies (Companies House)",sorted(companies,key=lambda x:x["CompanyName"]),CC,
 {"CompanyName":40,"UltimateOwner":36,"SIC1":42,"ConfidenceTier":34,"CompaniesHouseURL":46,"OwnershipCategory":30})

# ---------- GEOGRAPHY ----------
import re as _re
def parea(pc):
    m=_re.match(r'^([A-Z]{1,2})',(pc or "").upper().strip()); return m.group(1) if m else "?"
geo=collections.defaultdict(lambda: collections.Counter())
for p in practices:
    a=parea(p["Postcode"]); geo[a]["Total"]+=1
    k=p["OwnershipCategory"]
    if k.startswith("Independent"): geo[a]["Independent"]+=1
    elif k.startswith("Hakim"): geo[a]["Hakim Group"]+=1
    elif k.startswith("Group-owned"): geo[a]["Other group"]+=1
    else: geo[a]["National chain"]+=1
grows=[]
for a,c in sorted(geo.items(), key=lambda x:-x[1]["Total"]):
    grows.append({"PostcodeArea":a,"Total":c["Total"],"Independent":c["Independent"],
      "Hakim Group":c["Hakim Group"],"Other group":c["Other group"],"National chain":c["National chain"],
      "Independent %":round(c["Independent"]/c["Total"]*100,1)})
sheet("Geographic Summary",grows,["PostcodeArea","Total","Independent","Hakim Group","Other group","National chain","Independent %"],
      {"PostcodeArea":14,"Independent %":14})

# ---------- METHOD ----------
ws=wb.create_sheet("Method & Caveats")
ws.column_dimensions["A"].width=118
lines=[("METHOD & CAVEATS",TITLE),("",None),
("SOURCES",SUB),
("1. Companies House 'Basic Company Data' free bulk product, snapshot 01 Sep 2026 (5,689,367 companies).",None),
("2. Companies House Persons with Significant Control (PSC) bulk snapshot, 12 Sep 2026 (15.9m records).",None),
("3. NHS Organisation Data Service (ODS) live API - roles RO167 'Optical Site' and RO166 'Optical Headquarters'.",None),
("4. Practice websites and emails discovered by verified domain resolution (see caveats).",None),
("",None),
("HOW THE UNIVERSE WAS DEFINED",SUB),
("Practice level (England & Wales): every ACTIVE NHS 'Optical Site' record - 6,901 practices. This is the",None),
("most reliable list of actual trading optician premises, because it is where NHS sight tests are delivered.",None),
("",None),
("Company level (whole UK): all companies with SIC 47782 'Retail sale by opticians' (4,057), plus companies",None),
("under health SICs 86900/86220 whose name contains a strong optical term (2,342) - many optometry practices",None),
("register under the health SIC rather than the retail one. A further 806 weaker name matches are included and",None),
("flagged 'Tier C - review' so they can be checked rather than silently trusted.",None),
("",None),
("HOW OWNERSHIP WAS DETERMINED (the efficient method)",SUB),
("Each NHS practice links to an 'Optical Headquarters' record - the operating business. That HQ name was matched",None),
("to the Companies House register to get a company number, and the company number was then looked up in the PSC",None),
("data to find who actually controls it. Chains were additionally matched on brand name.",None),
("",None),
("Hakim Group specifically: their practices deliberately retain original local names, so name matching fails.",None),
("They were found via PSC control by HO2 Management Ltd and related vehicles registered at Unit 317, India Mill",None),
("Business Centre, Darwen BB3 1AE. This is the single most efficient identifier - one PSC query, no guesswork.",None),
("",None),
("CAVEATS - PLEASE READ",SUB),
("* England & Wales only at practice level. NHS ODS holds only 93 Scottish and 11 Northern Irish optical sites,",None),
("  because those nations run separate systems. Scotland and N. Ireland ARE covered in the company-level sheet.",None),
("* Roughly half of NHS optical HQs could not be matched to a limited company. These are overwhelmingly sole",None),
("  traders and partnerships, which have no Companies House entry - they are independent by definition, and are",None),
("  classified as such with legal form 'Unincorporated'.",None),
("* Company registered-office postcodes are often the accountant's address, not the practice. Practice postcodes",None),
("  in the practice sheets are the real premises.",None),
("* Website and email coverage is partial. Domains were found by generating candidates from the trading name and",None),
("  accepting a match only when the site resolved without cross-domain redirect AND its content confirmed it was",None),
("  an opticians. This is deliberately conservative: blank is better than wrong. Do not treat blanks as 'no website'.",None),
("* Ownership reflects the PSC snapshot date. Acquisitions completed after 12 Sep 2026 will not appear.",None),
("* 'Standalone' means the only practice under its ultimate owner in this dataset.",None),
("* A single brand can appear as more than one owner entity where branches are held in separate companies or",None),
("  partnerships (Armstrong & North is one example). Such brands may show as several smaller groups rather than",None),
("  one larger one, so the 'Owner League Table' understates a few regional chains.",None),
("* NHS Scotland publishes no equivalent optometry practice list (checked: 104 datasets on opendata.nhs.scot,",None),
("  none covering optometry), so Scotland cannot be added at practice level from an official open source.",None),
("* Practices that do no NHS work at all (a small number of purely private/boutique practices) will not appear",None),
("  in the NHS practice sheets, though their company will usually appear in the company-level sheet.",None)]
for i,(t,f) in enumerate(lines,1):
    c=ws.cell(i,1,t)
    if f: c.font=f
    c.alignment=Alignment(wrap_text=False,vertical="center")

wb.move_sheet("Headline Figures",offset=-10)
path="out/UK_Opticians_Ownership_Analysis.xlsx"
wb.save(path)
print("saved",path)
print("practices with website:",sum(1 for p in practices if p["Website"]),"email:",sum(1 for p in practices if p["Email"]))
